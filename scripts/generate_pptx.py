import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    BG_COLOR = RGBColor(8, 8, 8)
    SURFACE_COLOR = RGBColor(16, 16, 16)
    BORDER_COLOR = RGBColor(38, 38, 38)
    GREEN_COLOR = RGBColor(0, 255, 102)
    WHITE_COLOR = RGBColor(255, 255, 255)
    MUTED_COLOR = RGBColor(160, 160, 160)
    DIM_COLOR = RGBColor(100, 100, 100)
    AMBER_COLOR = RGBColor(255, 179, 0)
    RED_COLOR = RGBColor(255, 51, 68)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background() # no border
        return bg

    def add_masthead(slide, slide_num, category):
        # Top bar
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = RGBColor(12, 12, 12)
        top_bar.line.color.rgb = BORDER_COLOR
        top_bar.line.width = Pt(1)

        tf = top_bar.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = f"[POLICYARC]  //  {slide_num}  //  {category}"
        p.font.name = "Consolas"
        p.font.size = Pt(11)
        p.font.color.rgb = GREEN_COLOR
        p.font.bold = True

    # -------------------------------------------------------------
    # SLIDE 1: TITLE & EXECUTIVE HOOK
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    add_masthead(slide1, "SLIDE 01 / 05", "EXECUTIVE SUMMARY // CIRCLE ARC MICROGRANT")

    # Main Headline
    txBox = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.733), Inches(2.2))
    tf = txBox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "THE AUTONOMOUS AGENT IN THE FINANCIAL MACHINE"
    p0.font.name = "Consolas"
    p0.font.size = Pt(13)
    p0.font.color.rgb = GREEN_COLOR
    p0.font.bold = True

    p1 = tf.add_paragraph()
    p1.text = "POLICYARC: DETERMINISTIC AGENT SAFETY OS"
    p1.font.name = "Arial Black"
    p1.font.size = Pt(36)
    p1.font.color.rgb = WHITE_COLOR
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Enterprise-grade spending limits, supervisor co-signatures, and zero-volatility micro-settlement on Circle's Arc L1."
    p2.font.name = "Calibri"
    p2.font.size = Pt(17)
    p2.font.color.rgb = MUTED_COLOR
    p2.space_before = Pt(8)

    # Core Thesis Callout Box
    callout = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.7), Inches(11.733), Inches(1.1))
    callout.fill.solid()
    callout.fill.fore_color.rgb = RGBColor(16, 20, 16)
    callout.line.color.rgb = GREEN_COLOR
    callout.line.width = Pt(2)
    ctf = callout.text_frame
    ctf.margin_left = Inches(0.3)
    ctf.margin_top = Inches(0.15)
    cp = ctf.paragraphs[0]
    cp.text = "THE ARC ADVANTAGE: On Circle's Arc L1, USDC is the native gas asset. Transaction fees and settlement are denominated strictly in US dollars. PolicyArc pairs Arc's sub-second Malachite consensus (<800ms) with deterministic smart contract vaults to guarantee corporate treasury solvency."
    cp.font.name = "Consolas"
    cp.font.size = Pt(12)
    cp.font.color.rgb = WHITE_COLOR

    # 3 Metrics Cards
    card_w = Inches(3.7)
    card_h = Inches(1.8)
    card_gap = Inches(0.316)
    card_y = Inches(5.1)

    cards_data = [
        ("NATIVE GAS ASSET", "100% Native USDC", "Zero volatile tokens (ETH/SOL). Predictable fractions of a cent per execution."),
        ("BFT CONSENSUS", "< 800ms Finality", "Malachite consensus engine allows real-time agent micro-payments without lag."),
        ("SECURITY INVARIANTS", "7 / 7 Verified", "Formally audited under the Build-Harness engineering standard with 100% pass rate.")
    ]

    for i, (tag, val, desc) in enumerate(cards_data):
        cx = Inches(0.8) + i * (card_w + card_gap)
        c = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, card_y, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = SURFACE_COLOR
        c.line.color.rgb = BORDER_COLOR
        c.line.width = Pt(1)
        ctf = c.text_frame
        ctf.margin_left = Inches(0.2)
        ctf.margin_top = Inches(0.15)
        
        cp0 = ctf.paragraphs[0]
        cp0.text = tag
        cp0.font.name = "Consolas"
        cp0.font.size = Pt(10)
        cp0.font.color.rgb = DIM_COLOR
        
        cp1 = ctf.add_paragraph()
        cp1.text = val
        cp1.font.name = "Arial Black"
        cp1.font.size = Pt(17)
        cp1.font.color.rgb = GREEN_COLOR
        cp1.space_before = Pt(4)
        
        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.name = "Calibri"
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = MUTED_COLOR
        cp2.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 2: THE TWO FATAL CRISES
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_masthead(slide2, "SLIDE 02 / 05", "PROBLEM STATEMENT // STRUCTURAL BLOCKCHAIN FRICTION")

    txBox = slide2.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "THE TWO FATAL CRISES OF AUTONOMOUS AGENT FINANCE"
    p0.font.name = "Arial Black"
    p0.font.size = Pt(30)
    p0.font.color.rgb = WHITE_COLOR

    p1 = tf.add_paragraph()
    p1.text = "Why enterprises and developers cannot safely hand raw Web3 wallets to autonomous AI models:"
    p1.font.name = "Calibri"
    p1.font.size = Pt(16)
    p1.font.color.rgb = MUTED_COLOR
    p1.space_before = Pt(4)

    # 2 Big Crisis Cards
    col_w = Inches(5.7)
    col_h = Inches(4.2)
    
    # Crisis 1 Card
    c1 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.6), col_w, col_h)
    c1.fill.solid()
    c1.fill.fore_color.rgb = SURFACE_COLOR
    c1.line.color.rgb = RED_COLOR
    c1.line.width = Pt(2)
    ctf1 = c1.text_frame
    ctf1.margin_left = Inches(0.3)
    ctf1.margin_top = Inches(0.25)
    ctf1.word_wrap = True

    p = ctf1.paragraphs[0]
    p.text = "CRISIS 01: GAS PRICE VOLATILITY LOCKUP"
    p.font.name = "Consolas"
    p.font.size = Pt(13)
    p.font.color.rgb = RED_COLOR
    p.font.bold = True

    p = ctf1.add_paragraph()
    p.text = "Fluctuating Network Fees Stalling Workflows"
    p.font.name = "Arial Black"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE_COLOR
    p.space_before = Pt(6)

    bullets1 = [
        "The Two-Token Hurdle: Agents on Ethereum or Solana must hold volatile native assets (ETH, SOL) solely to pay execution gas.",
        "Sudden Fee Spikes: A network congestion spike or depleted gas balance immediately halts autonomous agent task loops mid-execution.",
        "Unpredictable Accounting: Corporate CFOs cannot budget or audit AI operations when infrastructure costs fluctuate 200% intra-day."
    ]
    for b in bullets1:
        p = ctf1.add_paragraph()
        p.text = "• " + b
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.color.rgb = MUTED_COLOR
        p.space_before = Pt(8)

    # Crisis 2 Card
    c2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.833), Inches(2.6), col_w, col_h)
    c2.fill.solid()
    c2.fill.fore_color.rgb = SURFACE_COLOR
    c2.line.color.rgb = RED_COLOR
    c2.line.width = Pt(2)
    ctf2 = c2.text_frame
    ctf2.margin_left = Inches(0.3)
    ctf2.margin_top = Inches(0.25)
    ctf2.word_wrap = True

    p = ctf2.paragraphs[0]
    p.text = "CRISIS 02: UNBOUNDED TREASURY DRAINAGE"
    p.font.name = "Consolas"
    p.font.size = Pt(13)
    p.font.color.rgb = RED_COLOR
    p.font.bold = True

    p = ctf2.add_paragraph()
    p.text = "Unconstrained Signing Autonomy & Attacks"
    p.font.name = "Arial Black"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE_COLOR
    p.space_before = Pt(6)

    bullets2 = [
        "Binary Key Permissions: Standard wallets give agents 100% total signing authority with zero spend limits or recipient controls.",
        "Prompt Injection Hijacking: A prompt injection or malicious API response can instruct the agent to transfer the entire corporate treasury into offshore pools.",
        "Runaway Inference Loops: Code bugs can trigger recursive infinite execution calls, draining thousands of dollars in minutes without warning."
    ]
    for b in bullets2:
        p = ctf2.add_paragraph()
        p.text = "• " + b
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.color.rgb = MUTED_COLOR
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 3: WHY CIRCLE ARC L1 IS LOAD-BEARING
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_masthead(slide3, "SLIDE 03 / 05", "SOLUTION ARCHITECTURE // THE ARC ECOSYSTEM FIT")

    txBox = slide3.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "WHY CIRCLE ARC L1 IS LOAD-BEARING FOR POLICYARC"
    p0.font.name = "Arial Black"
    p0.font.size = Pt(30)
    p0.font.color.rgb = WHITE_COLOR

    p1 = tf.add_paragraph()
    p1.text = "Arc is not just another EVM layer; its unique primitives make sovereign agent economics mathematically viable:"
    p1.font.name = "Calibri"
    p1.font.size = Pt(16)
    p1.font.color.rgb = MUTED_COLOR
    p1.space_before = Pt(4)

    # 3 Strategic Pillars
    pillar_w = Inches(3.7)
    pillar_h = Inches(4.3)
    pillar_gap = Inches(0.316)
    pillar_y = Inches(2.6)

    pillars = [
        ("PILLAR 01 // NATIVE GAS", "Native USDC Gas", [
            "Eliminates the two-token hurdle entirely.",
            "All transaction fees, gas accounting, and balances are denominated in dollars.",
            "Agents only hold and budget a single, stable currency.",
            "Gas costs are microscopic (~0.000120 USDC), enabling micro-payments down to $0.05."
        ]),
        ("PILLAR 02 // SPEED", "Malachite BFT (<800ms)", [
            "AI agents operate in high-speed algorithmic loops.",
            "Arc provides deterministic finality in under 800 milliseconds.",
            "No multi-minute block wait times or dropped nonces.",
            "Enables real-time agent-to-agent and agent-to-API payment settlement."
        ]),
        ("PILLAR 03 // GLOBAL EXPANSION", "Circle StableFX Ready", [
            "Native on-chain atomic stablecoin conversion engine.",
            "Agents can hold USDC and seamlessly route payments in EURC or global stablecoins.",
            "Enables global compute leasing across European and Asian localized clusters.",
            "Future-proof substrate for multinational enterprise agent workforces."
        ])
    ]

    for i, (tag, title, points) in enumerate(pillars):
        px = Inches(0.8) + i * (pillar_w + pillar_gap)
        p_card = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, px, pillar_y, pillar_w, pillar_h)
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = SURFACE_COLOR
        p_card.line.color.rgb = GREEN_COLOR if i == 0 else BORDER_COLOR
        p_card.line.width = Pt(1.5 if i == 0 else 1)
        ptf = p_card.text_frame
        ptf.margin_left = Inches(0.25)
        ptf.margin_top = Inches(0.2)
        ptf.word_wrap = True

        p = ptf.paragraphs[0]
        p.text = tag
        p.font.name = "Consolas"
        p.font.size = Pt(11)
        p.font.color.rgb = GREEN_COLOR
        p.font.bold = True

        p = ptf.add_paragraph()
        p.text = title
        p.font.name = "Arial Black"
        p.font.size = Pt(18)
        p.font.color.rgb = WHITE_COLOR
        p.space_before = Pt(6)

        for pt in points:
            p = ptf.add_paragraph()
            p.text = "• " + pt
            p.font.name = "Calibri"
            p.font.size = Pt(12)
            p.font.color.rgb = MUTED_COLOR
            p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 4: MECHANISM & INVARIANTS
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_masthead(slide4, "SLIDE 04 / 05", "MECHANISM & PROOF // BUILD-HARNESS VERIFICATION")

    txBox = slide4.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(1.3))
    tf = txBox.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "THE THREE-TIER SAFETY VALVE & 7 INVARIANTS"
    p0.font.name = "Arial Black"
    p0.font.size = Pt(30)
    p0.font.color.rgb = WHITE_COLOR

    p1 = tf.add_paragraph()
    p1.text = "PolicyVault implements deterministic on-chain governance backed by automated invariant testing:"
    p1.font.name = "Calibri"
    p1.font.size = Pt(15)
    p1.font.color.rgb = MUTED_COLOR

    # Left: The 3-Tier Mechanism
    m_box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.5), Inches(5.7), Inches(4.4))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = SURFACE_COLOR
    m_box.line.color.rgb = BORDER_COLOR
    mtf = m_box.text_frame
    mtf.margin_left = Inches(0.25)
    mtf.margin_top = Inches(0.2)
    mtf.word_wrap = True

    p = mtf.paragraphs[0]
    p.text = "ON-CHAIN GOVERNANCE ENGINE"
    p.font.name = "Consolas"
    p.font.size = Pt(11)
    p.font.color.rgb = GREEN_COLOR
    p.font.bold = True

    tiers = [
        ("Tier 1: Autonomous Spends (<= $2.00 USDC)", "Agent queries APIs or procures micro-tasks autonomously. Settles in <800ms with native USDC receipt."),
        ("Tier 2: Supervisor Co-Signature Gate (> $2.00)", "Preflight halts high-value spend. Requires an off-chain cryptographic ECDSA signature from human supervisor key."),
        ("Tier 3: Hard Ceilings ($10.00 Cap / $50 Daily)", "Non-negotiable invariants. Any transaction above $10 or daily spend above $50 immediately reverts on-chain."),
        ("Emergency Circuit Breaker", "Owner or supervisor can trigger fail-closed pause state instantly freezing all vault transfers.")
    ]
    for title, desc in tiers:
        p = mtf.add_paragraph()
        p.text = title
        p.font.name = "Arial Black"
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE_COLOR
        p.space_before = Pt(10)

        p = mtf.add_paragraph()
        p.text = desc
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.color.rgb = MUTED_COLOR

    # Right: Verified Invariants & Fresh Live Proofs
    i_box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.833), Inches(2.5), Inches(5.7), Inches(4.4))
    i_box.fill.solid()
    i_box.fill.fore_color.rgb = SURFACE_COLOR
    i_box.line.color.rgb = GREEN_COLOR
    i_box.line.width = Pt(1.5)
    itf = i_box.text_frame
    itf.margin_left = Inches(0.25)
    itf.margin_top = Inches(0.2)
    itf.word_wrap = True

    p = itf.paragraphs[0]
    p.text = "BUILD-HARNESS: 7/7 INVARIANTS PASS (100%)"
    p.font.name = "Consolas"
    p.font.size = Pt(11)
    p.font.color.rgb = GREEN_COLOR
    p.font.bold = True

    invs = [
        "INV-001: Per-Tx Spending Ceiling ($10.00 cap) [PASS]",
        "INV-002: 24h Rolling Budget Accumulator ($50.00) [PASS]",
        "INV-003: Supervisor ECDSA Dual-Signature Gateway [PASS]",
        "INV-004: Strict Destination Whitelist Registry [PASS]",
        "INV-005: Fail-Closed Emergency Circuit Breaker [PASS]",
        "INV-006: Reentrancy Protection via CEI Pattern [PASS]",
        "INV-007: Arc Native USDC Gas Predictability [PASS]"
    ]
    for inv in invs:
        p = itf.add_paragraph()
        p.text = "• " + inv
        p.font.name = "Consolas"
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE_COLOR
        p.space_before = Pt(4)

    p = itf.add_paragraph()
    p.text = "FRESH LIVE ARC TESTNET PROOFS:"
    p.font.name = "Consolas"
    p.font.size = Pt(11)
    p.font.color.rgb = AMBER_COLOR
    p.font.bold = True
    p.space_before = Pt(10)

    p = itf.add_paragraph()
    p.text = "• Tx 0xcc1e...f89b: -$0.25 USDC Autonomous Spend\n• Tx 0x0aae...78f9: -$3.50 USDC Supervisor Co-Signed\n• Verified on Arcscan Block #65,791,146"
    p.font.name = "Consolas"
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED_COLOR

    # -------------------------------------------------------------
    # SLIDE 5: ROADMAP & CIRCLE GRANT PIPELINE
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_masthead(slide5, "SLIDE 05 / 05", "ROADMAP // CIRCLE GRANT PROGRAM PIPELINE")

    txBox = slide5.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "FROM MICROGRANT TO ENTERPRISE SCALE"
    p0.font.name = "Arial Black"
    p0.font.size = Pt(30)
    p0.font.color.rgb = WHITE_COLOR

    p1 = tf.add_paragraph()
    p1.text = "The strategic roadmap positioning PolicyArc for the $10,000–$100,000 Circle Grant Program:"
    p1.font.name = "Calibri"
    p1.font.size = Pt(16)
    p1.font.color.rgb = MUTED_COLOR
    p1.space_before = Pt(4)

    # 3 Phase Cards
    phase_w = Inches(3.7)
    phase_h = Inches(3.2)
    phase_gap = Inches(0.316)
    phase_y = Inches(2.6)

    phases = [
        ("PHASE 01 (CURRENT)", "Arc Microgrant MVP", [
            "Complete working prototype on Arc Testnet.",
            "PolicyVault & MockUSDC contracts live.",
            "7 Formally verified security invariants.",
            "Interactive operational console & docs."
        ]),
        ("PHASE 02 (Q4 2026)", "Circle StableFX Integration", [
            "Integrate Circle's native StableFX engine.",
            "Atomic multi-stablecoin swaps (USDC/EURC).",
            "Cross-border agent payroll settlement.",
            "Localized compute cluster leasing."
        ]),
        ("PHASE 03 (Q1 2027)", "Departmental Pools & SDKs", [
            "Hierarchical multi-agent corporate pools.",
            "Sub-agent micro-budget delegation.",
            "First-party plugins for LangChain & CrewAI.",
            "Audit & compliance export for corporate ERPs."
        ])
    ]

    for i, (tag, title, points) in enumerate(phases):
        px = Inches(0.8) + i * (phase_w + phase_gap)
        c = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, px, phase_y, phase_w, phase_h)
        c.fill.solid()
        c.fill.fore_color.rgb = SURFACE_COLOR
        c.line.color.rgb = GREEN_COLOR if i == 0 else BORDER_COLOR
        c.line.width = Pt(1)
        ctf = c.text_frame
        ctf.margin_left = Inches(0.2)
        ctf.margin_top = Inches(0.15)
        ctf.word_wrap = True

        p = ctf.paragraphs[0]
        p.text = tag
        p.font.name = "Consolas"
        p.font.size = Pt(10)
        p.font.color.rgb = GREEN_COLOR
        p.font.bold = True

        p = ctf.add_paragraph()
        p.text = title
        p.font.name = "Arial Black"
        p.font.size = Pt(16)
        p.font.color.rgb = WHITE_COLOR
        p.space_before = Pt(4)

        for pt in points:
            p = ctf.add_paragraph()
            p.text = "• " + pt
            p.font.name = "Calibri"
            p.font.size = Pt(11)
            p.font.color.rgb = MUTED_COLOR
            p.space_before = Pt(6)

    # Bottom Submission Links Box
    sub_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.733), Inches(1.0))
    sub_box.fill.solid()
    sub_box.fill.fore_color.rgb = RGBColor(12, 16, 12)
    sub_box.line.color.rgb = GREEN_COLOR
    sub_box.line.width = Pt(1.5)
    stf = sub_box.text_frame
    stf.margin_left = Inches(0.25)
    stf.margin_top = Inches(0.12)
    stf.word_wrap = True

    p = stf.paragraphs[0]
    p.text = "OFFICIAL SUBMISSION LINKS & ARTIFACTS:"
    p.font.name = "Consolas"
    p.font.size = Pt(11)
    p.font.color.rgb = GREEN_COLOR
    p.font.bold = True

    p = stf.add_paragraph()
    p.text = "• Live Console: https://policyarc.vercel.app   • Docs: https://policyarc.vercel.app/docs\n• GitHub: https://github.com/0xNexuz/policy-arc   • Contract: 0xb9176558920B53e9542AD75eDC9Bfa68de3A89C0"
    p.font.name = "Consolas"
    p.font.size = Pt(10)
    p.font.color.rgb = WHITE_COLOR

    output_path = os.path.join(os.path.dirname(__file__), "../PolicyArc_Pitch_Deck.pptx")
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    create_deck()
